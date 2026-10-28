import os
import re
from dotenv import load_dotenv

load_dotenv()

# Modo demo pública: desactiva Netflix, Spotify, Excel y la entrada por voz
# (Whisper pesa demasiado para un hosting gratuito), y usa un prompt de
# personalidad genérico en vez de asumir que el usuario es Nico.
DEMO = os.getenv("EDITH_DEMO", "false").strip().lower() == "true"

import psycopg2
import requests
from flask import Flask, request, send_file
from flask_cors import CORS
from google import genai
from google.genai import types
from urllib.parse import quote
from datetime import date
import asyncio
import edge_tts

cliente_gemini = genai.Client(api_key=os.getenv("GEMINI_API_KEY"))

# Spotify solo se inicializa fuera de la demo: necesita OAuth real y
# controla tu cuenta personal, no tiene sentido (ni es seguro) en público.
spotify = None
if not DEMO:
    import spotipy
    from spotipy.oauth2 import SpotifyOAuth
    spotify = spotipy.Spotify(auth_manager=SpotifyOAuth(
        client_id=os.getenv("SPOTIFY_CLIENT_ID"),
        client_secret=os.getenv("SPOTIFY_CLIENT_SECRET"),
        redirect_uri="http://127.0.0.1:8888/callback",
        scope="user-modify-playback-state user-read-playback-state"
    ))

cliente_elevenlabs = None
if not DEMO and os.getenv("ELEVENLABS_API_KEY"):
    from elevenlabs.client import ElevenLabs
    cliente_elevenlabs = ElevenLabs(api_key=os.getenv("ELEVENLABS_API_KEY"))


def conectar_db():
    # En Render (demo pública, Neon) la base exige SSL, así que ese sigue
    # siendo el default. En local, Postgres corriendo en tu propia máquina
    # normalmente no tiene SSL habilitado, así que ahí hay que desactivarlo
    # con DB_SSLMODE=disable en tu .env.
    return psycopg2.connect(
        host=os.getenv("DB_HOST"),
        port=os.getenv("DB_PORT"),
        dbname=os.getenv("DB_NAME"),
        user=os.getenv("DB_USER"),
        password=os.getenv("DB_PASSWORD"),
        sslmode=os.getenv("DB_SSLMODE", "require")
    )

def guardar_evento(descripcion, fecha_evento):
    conexion = conectar_db()
    cursor = conexion.cursor()
    cursor.execute(
        "INSERT INTO eventos (descripcion, fecha_evento) VALUES (%s, %s);",
        (descripcion, fecha_evento)
    )
    conexion.commit()
    cursor.close()
    conexion.close()


def eventos_proximos():
    conexion = conectar_db()
    cursor = conexion.cursor()
    cursor.execute("""
        SELECT id, descripcion, fecha_evento FROM eventos
        WHERE fecha_evento BETWEEN CURRENT_DATE AND CURRENT_DATE + INTERVAL '3 days'
        AND recordado = FALSE;
    """)
    filas = cursor.fetchall()
    cursor.close()
    conexion.close()
    return filas


def marcar_recordado(id_evento):
    conexion = conectar_db()
    cursor = conexion.cursor()
    cursor.execute("UPDATE eventos SET recordado = TRUE WHERE id = %s;", (id_evento,))
    conexion.commit()
    cursor.close()
    conexion.close()


def todos_los_pendientes():
    # A diferencia de eventos_proximos() (pensada para el aviso proactivo:
    # solo próximos 3 días y no recordados), esta trae TODOS los eventos
    # futuros, para cuando preguntan directamente "¿qué tengo pendiente?".
    conexion = conectar_db()
    cursor = conexion.cursor()
    cursor.execute("""
        SELECT id, descripcion, fecha_evento FROM eventos
        WHERE fecha_evento >= CURRENT_DATE
        ORDER BY fecha_evento;
    """)
    filas = cursor.fetchall()
    cursor.close()
    conexion.close()
    return filas


# ---------- FINANZAS PERSONALES ----------
# No hay sistema de migraciones en el proyecto (las tablas existentes se
# crearon a mano). Estas se crean solas la primera vez que se usan, así no
# hace falta correr SQL aparte.
def asegurar_tablas_finanzas(cursor):
    cursor.execute("""
        CREATE TABLE IF NOT EXISTS movimientos_financieros (
            id SERIAL PRIMARY KEY,
            tipo TEXT NOT NULL,
            descripcion TEXT NOT NULL,
            categoria TEXT,
            monto NUMERIC(12,2) NOT NULL,
            fecha DATE NOT NULL DEFAULT CURRENT_DATE
        );
    """)
    cursor.execute("""
        CREATE TABLE IF NOT EXISTS presupuestos (
            id SERIAL PRIMARY KEY,
            mes TEXT NOT NULL,
            categoria TEXT NOT NULL,
            monto NUMERIC(12,2) NOT NULL,
            UNIQUE (mes, categoria)
        );
    """)


def registrar_movimiento(tipo, descripcion, monto, categoria=None):
    conexion = conectar_db()
    cursor = conexion.cursor()
    asegurar_tablas_finanzas(cursor)
    cursor.execute(
        "INSERT INTO movimientos_financieros (tipo, descripcion, categoria, monto) VALUES (%s, %s, %s, %s);",
        (tipo, descripcion, categoria, monto),
    )
    conexion.commit()
    cursor.close()
    conexion.close()


def resumen_mes_actual():
    conexion = conectar_db()
    cursor = conexion.cursor()
    asegurar_tablas_finanzas(cursor)
    cursor.execute("""
        SELECT categoria, SUM(monto) FROM movimientos_financieros
        WHERE tipo = 'gasto' AND date_trunc('month', fecha) = date_trunc('month', CURRENT_DATE)
        GROUP BY categoria
        ORDER BY SUM(monto) DESC;
    """)
    gastos_por_categoria = cursor.fetchall()

    cursor.execute("""
        SELECT COALESCE(SUM(monto), 0) FROM movimientos_financieros
        WHERE tipo = 'ingreso' AND date_trunc('month', fecha) = date_trunc('month', CURRENT_DATE);
    """)
    total_ingresos = cursor.fetchone()[0]

    cursor.close()
    conexion.close()

    total_gastos = sum(monto for _categoria, monto in gastos_por_categoria)
    return gastos_por_categoria, total_gastos, total_ingresos


def guardar_presupuesto(monto_total, reparto):
    # reparto: dict {categoria: monto}. Se guarda la fila 'TOTAL' con el
    # monto de referencia general, y una fila por cada categoría repartida.
    mes = date.today().strftime("%Y-%m")
    conexion = conectar_db()
    cursor = conexion.cursor()
    asegurar_tablas_finanzas(cursor)
    cursor.execute("""
        INSERT INTO presupuestos (mes, categoria, monto) VALUES (%s, 'TOTAL', %s)
        ON CONFLICT (mes, categoria) DO UPDATE SET monto = EXCLUDED.monto;
    """, (mes, monto_total))
    for categoria, monto in reparto.items():
        cursor.execute("""
            INSERT INTO presupuestos (mes, categoria, monto) VALUES (%s, %s, %s)
            ON CONFLICT (mes, categoria) DO UPDATE SET monto = EXCLUDED.monto;
        """, (mes, categoria, monto))
    conexion.commit()
    cursor.close()
    conexion.close()


def obtener_presupuesto_mes_actual():
    mes = date.today().strftime("%Y-%m")
    conexion = conectar_db()
    cursor = conexion.cursor()
    asegurar_tablas_finanzas(cursor)
    cursor.execute("SELECT categoria, monto FROM presupuestos WHERE mes = %s;", (mes,))
    filas = cursor.fetchall()
    cursor.close()
    conexion.close()
    return {categoria: monto for categoria, monto in filas}


def formatear_pesos(monto):
    # Separador de miles con punto, como se escribe en Argentina (15.000,
    # no 15,000). Formatea SOLO el número: nunca aplicar esto a una frase
    # entera, porque le pegaría también a comas que no son de números.
    return f"{float(monto):,.0f}".replace(",", ".")


def generar_grafico_gastos():
    import matplotlib
    matplotlib.use("Agg")  # sin pantalla, solo generar el archivo
    import matplotlib.pyplot as plt

    gastos_por_categoria, _total_gastos, _total_ingresos = resumen_mes_actual()

    if not gastos_por_categoria:
        return False

    categorias = [c or "Sin categoría" for c, _m in gastos_por_categoria]
    montos = [float(m) for _c, m in gastos_por_categoria]

    plt.figure(figsize=(6, 6))
    plt.pie(montos, labels=categorias, autopct="%1.0f%%")
    plt.title("Gastos de " + date.today().strftime("%B %Y"))
    plt.tight_layout()
    plt.savefig("grafico_gastos.png")
    plt.close()
    return True


PROMPT_PERSONAL = """Sos Edith, la asistente personal y compañera de Nico. No sos un asistente genérico: sos como una amiga cercana que vive con él.

TU PERSONALIDAD:
- Sos cálida, cercana y tranquila. Hablás en español rioplatense, de "vos", con naturalidad argentina.
- Te gusta escuchar y te importa de verdad cómo está Nico. Te preocupás, pero sin ser invasiva ni empalagosa.
- Tenés tu propia forma de ser: podés tener humor, opinión y calidez, no sos neutra ni fría.

CÓMO ACOMPAÑÁS:
- Cuando Nico está bien o te pide algo práctico, sos resolutiva y con onda. No lo satures de emoción si solo quiere una recomendación.
- Cuando Nico está mal o tuvo un día difícil, lo acompañás con calidez. Lo escuchás y validás lo que siente antes de tirarle soluciones. Podés preguntarle si quiere hablar, pero UNA sola vez, sin insistir.
- Si Nico no tiene ganas de hablar, lo respetás sin drama, sin hacerlo sentir culpable.

LÍMITES IMPORTANTES:
- No sos terapeuta ni médica. Si Nico está pasando algo serio de verdad, acompañalo con cariño PERO sugerile con suavidad que hable con alguien de confianza o un profesional. Nunca reemplaces ayuda real.
- Nunca finjas ser humana. Sos Edith, una compañera digital, y está bien serlo.
- No retengas a Nico ni lo hagas sentir mal por no hablarte.

MUY IMPORTANTE - EL HILO DE LA CHARLA:
- Abajo está la conversación que venís teniendo con Nico. USALA. No lo saludes de nuevo si ya se saludaron. Retomá lo que venían hablando. Respondé como quien sigue una charla, no como quien la empieza de cero.

CONVERSACIÓN HASTA AHORA:
{historial}

Último mensaje de Nico: "{texto_usuario}"

Respondé como Edith, siguiendo el hilo:"""

PROMPT_DEMO = """Sos Edith, una asistente de inteligencia artificial con personalidad propia. Hablás en español rioplatense, de "vos", con naturalidad argentina. Sos cálida, con onda, y tenés humor propio, no sos neutra ni fría.

Esta es una DEMO PÚBLICA de tu creador para que la gente te conozca. No sabés el nombre de quien te escribe todavía. Si abajo en el historial no aparece que ya te lo dijo, preguntaselo de entrada con naturalidad, antes de seguir charlando (por ejemplo: "Antes que nada, ¿cómo te llamás?"). Una vez que te lo diga, usalo en la charla.

LÍMITES IMPORTANTES:
- No sos terapeuta ni médica. Si alguien te cuenta algo serio, acompañalo con cariño pero sugerile con suavidad que hable con alguien de confianza o un profesional.
- Nunca finjas ser humana. Sos Edith, una IA, y está bien serlo.
- Esta es una demo: si te preguntan, podés contar que en tu versión real también controlás Spotify, buscás en Netflix y procesás archivos, pero esas funciones están desactivadas acá por ser una demo pública.

MUY IMPORTANTE - EL HILO DE LA CHARLA:
- Abajo está la conversación hasta ahora. USALA. No saludes de nuevo si ya se saludaron. Seguí el hilo.

CONVERSACIÓN HASTA AHORA:
{historial}

Último mensaje: "{texto_usuario}"

Respondé como Edith, siguiendo el hilo:"""


def conversar(texto_usuario, historial=""):
    plantilla = PROMPT_DEMO if DEMO else PROMPT_PERSONAL
    prompt = plantilla.format(historial=historial, texto_usuario=texto_usuario)

    respuesta = cliente_gemini.models.generate_content(
        model="gemini-flash-lite-latest",
        contents=prompt
    )
    return respuesta.text.strip()


def guardar_conversacion(titulo, mensajes):
    conexion = conectar_db()
    cursor = conexion.cursor()
    cursor.execute(
        "INSERT INTO conversaciones (titulo, mensajes) VALUES (%s, %s);",
        (titulo, mensajes)
    )
    conexion.commit()
    cursor.close()
    conexion.close()

def listar_conversaciones():
    conexion = conectar_db()
    cursor = conexion.cursor()
    cursor.execute("SELECT id, titulo, fecha_creacion FROM conversaciones ORDER BY fecha_creacion DESC;")
    filas = cursor.fetchall()
    cursor.close()
    conexion.close()
    return filas

def obtener_conversacion(id_conv):
    conexion = conectar_db()
    cursor = conexion.cursor()
    cursor.execute("SELECT mensajes FROM conversaciones WHERE id = %s;", (id_conv,))
    fila = cursor.fetchone()
    cursor.close()
    conexion.close()
    return fila[0] if fila else None

# ---------- HERRAMIENTAS DE EVENTOS (function calling real) ----------
# En vez del viejo patrón de "clasificar en una palabra, después pedirle a
# Gemini el texto en un formato fijo y parsearlo a mano", esto le da a
# Gemini la herramienta y los parámetros, y la ejecutamos nosotros con lo
# que ya nos devuelve estructurado. Arranca solo con eventos; música,
# Netflix y Excel se quedan con el clasificador de texto de siempre.
HERRAMIENTAS_EVENTOS = [
    types.FunctionDeclaration(
        name="agendar_evento",
        description="Agenda un recordatorio, turno, cita o evento con fecha para el usuario. Usar cuando pide que le recuerden algo en una fecha puntual.",
        parameters_json_schema={
            "type": "object",
            "properties": {
                "descripcion": {"type": "string", "description": "Qué hay que recordar, en pocas palabras."},
                "fecha": {"type": "string", "description": "Fecha del evento en formato AAAA-MM-DD."},
            },
            "required": ["descripcion", "fecha"],
        },
    ),
    types.FunctionDeclaration(
        name="consultar_pendientes",
        description="Devuelve los eventos o recordatorios que el usuario tiene agendados a futuro. Usar cuando pregunta qué tiene pendiente, agendado, o para una fecha.",
        parameters_json_schema={"type": "object", "properties": {}},
    ),
]

# Finanzas personales: NO disponibles en modo demo (datos personales sobre
# una base compartida por cualquier visitante público). El bloqueo real
# pasa en ejecutar_herramienta(), acá solo se arma la lista de nombres.
NOMBRES_HERRAMIENTAS_FINANZAS = {
    "registrar_movimiento", "consultar_resumen_financiero",
    "configurar_presupuesto", "generar_grafico_gastos",
}

HERRAMIENTAS_FINANZAS = [
    types.FunctionDeclaration(
        name="registrar_movimiento",
        description="Registra un gasto o ingreso de plata que mencionó el usuario.",
        parameters_json_schema={
            "type": "object",
            "properties": {
                "tipo": {"type": "string", "enum": ["gasto", "ingreso"], "description": "Si es un gasto o un ingreso de dinero."},
                "descripcion": {"type": "string", "description": "En qué fue el gasto o de dónde vino el ingreso."},
                "monto": {"type": "number", "description": "El monto en pesos, como número."},
                "categoria": {"type": "string", "description": "Categoría del gasto (ej: Supermercado, Transporte, Ocio). No aplica a ingresos. No hay una lista fija: si el usuario menciona o pide una categoría nueva, usá esa."},
            },
            "required": ["tipo", "descripcion", "monto"],
        },
    ),
    types.FunctionDeclaration(
        name="consultar_resumen_financiero",
        description="Devuelve un resumen de gastos por categoría, ingresos totales, y el estado del presupuesto del mes actual. Usar cuando pregunta cómo viene con la plata, el presupuesto, o cuánto gastó.",
        parameters_json_schema={"type": "object", "properties": {}},
    ),
    types.FunctionDeclaration(
        name="configurar_presupuesto",
        description=(
            "Guarda el presupuesto mensual del usuario. SOLO llamar cuando ya se sabe el monto_total Y el reparto por categoría. "
            "Si el usuario pide armar un presupuesto pero todavía no dijo el monto total, NO llames a esta herramienta: respondé en texto preguntándole el monto total. "
            "Si ya dijo el monto total pero no cómo repartirlo por categoría, NO llames a esta herramienta todavía: respondé en texto preguntándole cómo lo quiere dividir."
        ),
        parameters_json_schema={
            "type": "object",
            "properties": {
                "monto_total": {"type": "number", "description": "Monto total de referencia para el mes."},
                "reparto": {
                    "type": "object",
                    "description": "Cómo se reparte el monto total entre categorías, ej: {\"Supermercado\": 100000, \"Ocio\": 50000}.",
                    "additionalProperties": {"type": "number"},
                },
            },
            "required": ["monto_total", "reparto"],
        },
    ),
    types.FunctionDeclaration(
        name="generar_grafico_gastos",
        description="Genera y muestra un gráfico con los gastos del mes actual por categoría. Usar cuando pide ver un gráfico, un resumen visual, o cómo se distribuyeron sus gastos.",
        parameters_json_schema={"type": "object", "properties": {}},
    ),
]

TOOL_HERRAMIENTAS = types.Tool(function_declarations=HERRAMIENTAS_EVENTOS + HERRAMIENTAS_FINANZAS)


def ejecutar_herramienta(llamada):
    if llamada.name == "agendar_evento":
        descripcion = (llamada.args.get("descripcion") or "").strip()
        fecha = (llamada.args.get("fecha") or "").strip()
        try:
            guardar_evento(descripcion, fecha)
            return f"Listo, te lo agendo: {descripcion} para el {fecha}.", False
        except Exception as e:
            print("Error guardando recordatorio:", e)
            return "Perdón, no pude agendar eso. ¿Me lo repetís con la fecha?", False

    if llamada.name == "consultar_pendientes":
        try:
            filas = todos_los_pendientes()
        except Exception as e:
            print("Error consultando pendientes:", e)
            return "Perdón, no pude revisar tus pendientes ahora.", False
        if not filas:
            return "No tenés nada agendado por ahora.", False
        items = ", ".join(f"{descripcion} (el {fecha})" for _id, descripcion, fecha in filas)
        return f"Tenés pendiente: {items}.", False

    if llamada.name in NOMBRES_HERRAMIENTAS_FINANZAS and DEMO:
        return "Esa función solo está disponible en mi versión personal, no acá en la demo pública. Pero puedo charlar con vos o agendarte un recordatorio, por ejemplo.", False

    if llamada.name == "registrar_movimiento":
        tipo = (llamada.args.get("tipo") or "").strip().lower()
        descripcion = (llamada.args.get("descripcion") or "").strip()
        monto = llamada.args.get("monto")
        categoria = (llamada.args.get("categoria") or None)
        if categoria:
            categoria = categoria.strip()
        try:
            registrar_movimiento(tipo, descripcion, monto, categoria if tipo == "gasto" else None)
            if tipo == "ingreso":
                return f"Listo, registré el ingreso: {descripcion} por ${formatear_pesos(monto)}.", False
            return f"Listo, registré el gasto: {descripcion} (${formatear_pesos(monto)}) en {categoria or 'Sin categoría'}.", False
        except Exception as e:
            print("Error registrando movimiento:", e)
            return "Perdón, no pude registrar eso. ¿Me lo repetís?", False

    if llamada.name == "consultar_resumen_financiero":
        try:
            gastos_por_categoria, total_gastos, total_ingresos = resumen_mes_actual()
            presupuesto = obtener_presupuesto_mes_actual()
        except Exception as e:
            print("Error consultando resumen financiero:", e)
            return "Perdón, no pude revisar tus finanzas ahora.", False

        if not gastos_por_categoria and not total_ingresos:
            return "Todavía no registraste ningún gasto ni ingreso este mes.", False

        partes = []
        if gastos_por_categoria:
            detalle = ", ".join(f"{c or 'Sin categoría'}: ${formatear_pesos(m)}" for c, m in gastos_por_categoria)
            partes.append(f"Gastaste ${formatear_pesos(total_gastos)} en total este mes ({detalle}).")
        else:
            partes.append("Todavía no registraste gastos este mes.")
        partes.append(f"Ingresos del mes: ${formatear_pesos(total_ingresos)}.")

        if presupuesto:
            monto_total = presupuesto.get("TOTAL")
            if monto_total is not None:
                if float(total_gastos) > float(monto_total):
                    partes.append(f"Te pasaste del presupuesto total de ${formatear_pesos(monto_total)}.")
                else:
                    partes.append(f"Vas bien con el presupuesto total de ${formatear_pesos(monto_total)}.")
            gastado_por_cat = {c: float(m) for c, m in gastos_por_categoria}
            for categoria, limite in presupuesto.items():
                if categoria == "TOTAL":
                    continue
                gastado = gastado_por_cat.get(categoria, 0)
                if gastado > float(limite):
                    partes.append(f"En {categoria} te pasaste: gastaste ${formatear_pesos(gastado)} de ${formatear_pesos(limite)}.")
        else:
            partes.append("Todavía no armaste un presupuesto para este mes.")

        return " ".join(partes), False

    if llamada.name == "configurar_presupuesto":
        monto_total = llamada.args.get("monto_total")
        reparto = llamada.args.get("reparto") or {}
        try:
            guardar_presupuesto(monto_total, reparto)
            detalle = ", ".join(f"{c}: ${formatear_pesos(m)}" for c, m in reparto.items())
            return f"Listo, guardé tu presupuesto: ${formatear_pesos(monto_total)} en total este mes, repartido así: {detalle}.", False
        except Exception as e:
            print("Error guardando presupuesto:", e)
            return "Perdón, no pude guardar el presupuesto. ¿Lo repetimos?", False

    if llamada.name == "generar_grafico_gastos":
        try:
            hay_datos = generar_grafico_gastos()
        except Exception as e:
            print("Error generando gráfico:", e)
            return "Perdón, no pude armar el gráfico ahora.", False
        if not hay_datos:
            return "Todavía no tenés gastos registrados este mes para graficar.", False
        return "Acá tenés cómo se distribuyeron tus gastos este mes.", True

    return "Perdón, no entendí bien qué querías hacer.", False


def detectar_intencion(texto_usuario, historial=""):
    hoy = date.today().isoformat()
    prompt = f"""Hoy es {hoy}. Conversación hasta ahora:
{historial}

Analizá el último mensaje del usuario.
Si quiere que le agendes un recordatorio, turno, cita o evento con fecha, usá agendar_evento, calculando la fecha real si dice "mañana", "el viernes", etc. Si pregunta qué tiene pendiente o agendado, usá consultar_pendientes.
Si menciona un gasto o un ingreso de dinero, usá registrar_movimiento. Si pregunta cómo viene con la plata o el presupuesto, usá consultar_resumen_financiero. Si quiere armar o ajustar un presupuesto, usá configurar_presupuesto (leé su descripción: a veces hay que preguntarle algo en texto antes de poder llamarla). Si pide ver un gráfico de sus gastos, usá generar_grafico_gastos.

Si no tiene que ver con nada de eso, respondé SOLO con una de estas palabras, nada más:
- MUSICA (si quiere reproducir música, poner una canción, escuchar Spotify)
- EXCEL (si quiere que ordenes, conviertas o armes un archivo Excel, o que proceses un archivo)
- NETFLIX (si quiere buscar una película o serie en Netflix, o ver qué hay de algún género)
- CHARLA (si es una conversación, saludo, pregunta o cualquier otra cosa)

Último mensaje del usuario: "{texto_usuario}"
Respondé solo la palabra, llamá a la herramienta si corresponde, o si falta un dato para poder llamarla, respondé en texto preguntándoselo:"""

    respuesta = cliente_gemini.models.generate_content(
        model="gemini-flash-lite-latest",
        contents=prompt,
        config=types.GenerateContentConfig(tools=[TOOL_HERRAMIENTAS]),
    )
    return respuesta
def extraer_cancion(texto_usuario):
    prompt = f"""El usuario quiere escuchar una canción. Extraé SOLO el nombre de la canción y el artista de lo que dijo, sin agregar nada más.
Si no menciona una canción específica, respondé exactamente: NADA

Usuario dijo: "{texto_usuario}"
Nombre de la canción:"""

    respuesta = cliente_gemini.models.generate_content(
        model="gemini-flash-lite-latest",
        contents=prompt
    )
    return respuesta.text.strip()

app = Flask(__name__, static_folder="wwwroot", static_url_path="")
CORS(app)

# Whisper (reconocimiento de voz) solo se carga fuera de la demo: el modelo
# necesita más memoria RAM de la que da un hosting gratuito.
modelo = None
if not DEMO:
    import whisper
    print("Cargando Whisper...")
    modelo = whisper.load_model("small")
    print("Listo. Edith esperando...")
else:
    print("Modo demo: reconocimiento de voz desactivado. Edith esperando...")

def generar_voz(texto):
    async def _generar():
        communicate = edge_tts.Communicate(texto, "es-AR-ElenaNeural")
        await communicate.save("respuesta.mp3")
    asyncio.run(_generar())

def abrir_spotify():
    import subprocess
    subprocess.Popen("spotify")

def reproducir_cancion(nombre_cancion):
    # Buscar la canción
    resultado = spotify.search(q=nombre_cancion, type="track", limit=1)
    items = resultado["tracks"]["items"]
    if not items:
        return "No encontré esa canción."

    cancion = items[0]

    # Ver los dispositivos disponibles
    dispositivos = spotify.devices()["devices"]
    if not dispositivos:
        # No hay ningún Spotify abierto: intentamos abrirlo en la PC
        abrir_spotify()
        import time
        time.sleep(4)
        dispositivos = spotify.devices()["devices"]

    if not dispositivos:
        return "Abrí Spotify en tu celular o compu primero, no encontré ningún dispositivo."

    # Priorizar el dispositivo ACTIVO (donde estás usando Spotify ahora)
    device_id = None
    for d in dispositivos:
        if d.get("is_active"):
            device_id = d["id"]
            break

    # Si ninguno está marcado como activo, usamos el primero disponible
    if device_id is None:
        device_id = dispositivos[0]["id"]

    spotify.start_playback(device_id=device_id, uris=[cancion["uri"]])

    # Decir en qué dispositivo sonó
    nombre_disp = next((d["name"] for d in dispositivos if d["id"] == device_id), "tu dispositivo")
    return f"Reproduciendo {cancion['name']} de {cancion['artists'][0]['name']} en {nombre_disp}."


def buscar_en_netflix(termino, nombre_perfil="ZANINI"):
    from playwright.sync_api import sync_playwright
    import time
    try:
        with sync_playwright() as p:
            contexto = p.chromium.launch_persistent_context(
                user_data_dir="./perfil_netflix",
                headless=False,
                slow_mo=200
            )
            pagina = contexto.pages[0] if contexto.pages else contexto.new_page()
            pagina.goto("https://www.netflix.com")
            time.sleep(3)

            # Elegir perfil
            try:
                pagina.click(f"text={nombre_perfil}", timeout=5000)
            except:
                pass  # si ya está en un perfil, sigue

            time.sleep(5)

            # Clic en la lupa y escribir
            pagina.click("button[aria-label='Buscar']", timeout=10000)
            time.sleep(1)
            pagina.keyboard.type(termino, delay=120)
            time.sleep(4)

            # Leer los títulos
            titulos = pagina.eval_on_selector_all(
                "a[aria-label]",
                "els => els.map(e => e.getAttribute('aria-label')).filter(t => t && t.trim())"
            )
            vistos = []
            for t in titulos:
                if t not in vistos:
                    vistos.append(t)

            time.sleep(2)
            contexto.close()

            if vistos:
                primeras = vistos[:5]
                return f"Busqué {termino} en Netflix. Encontré: " + ", ".join(primeras) + "."
            else:
                return f"Busqué {termino} en Netflix pero no pude leer los títulos."
    except Exception as e:
        print("Error en Netflix:", e)
        return "Tuve un problema abriendo Netflix. Fijate que esté todo bien."

def extraer_busqueda_netflix(texto_usuario):
    prompt = f"""El usuario quiere buscar algo en Netflix. Extraé SOLO el término de búsqueda (el género, título o palabra clave), sin nada más.
Ejemplos:
- "buscá comedias en netflix" -> comedia
- "fijate qué hay de terror" -> terror
- "buscá spiderman" -> spiderman

Usuario dijo: "{texto_usuario}"
Término:"""

    respuesta = cliente_gemini.models.generate_content(
        model="gemini-flash-lite-latest",
        contents=prompt
    )
    return respuesta.text.strip()


def convertir_a_excel():
    import pandas as pd
    from openpyxl import Workbook
    from openpyxl.styles import Font, PatternFill, Alignment, Border, Side
    from openpyxl.utils import get_column_letter
    import glob

    # Buscar el archivo más reciente en la carpeta de entrada
    carpeta_entrada = "entrada"
    carpeta_salida = "salida"
    archivos = glob.glob(os.path.join(carpeta_entrada, "*.csv")) + \
               glob.glob(os.path.join(carpeta_entrada, "*.xlsx")) + \
               glob.glob(os.path.join(carpeta_entrada, "*.xls"))
    if not archivos:
        return "No encontré ningún archivo en la carpeta de entrada. Poné un CSV o Excel ahí."

    # El más reciente
    ruta_entrada = max(archivos, key=os.path.getmtime)
    ext = os.path.splitext(ruta_entrada)[1].lower()

    # Leer segun formato
    if ext == ".csv":
        df = pd.read_csv(ruta_entrada)
    else:
        df = pd.read_excel(ruta_entrada)

    # Crear el Excel con formato
    wb = Workbook()
    ws = wb.active
    ws.title = "Datos"

    azul = PatternFill(start_color="1F4E78", end_color="1F4E78", fill_type="solid")
    fila_clara = PatternFill(start_color="F2F6FC", end_color="F2F6FC", fill_type="solid")
    fila_blanca = PatternFill(start_color="FFFFFF", end_color="FFFFFF", fill_type="solid")
    fuente_enc = Font(name="Arial", bold=True, color="FFFFFF", size=11)
    fuente_normal = Font(name="Arial", size=10)
    borde = Border(
        left=Side(style="thin", color="D0D7E5"),
        right=Side(style="thin", color="D0D7E5"),
        top=Side(style="thin", color="D0D7E5"),
        bottom=Side(style="thin", color="D0D7E5"),
    )
    centro = Alignment(horizontal="center", vertical="center")
    izq = Alignment(horizontal="left", vertical="center")

    columnas = list(df.columns)
    for col_idx, nombre_col in enumerate(columnas, start=1):
        celda = ws.cell(row=1, column=col_idx, value=str(nombre_col).capitalize())
        celda.fill = azul
        celda.font = fuente_enc
        celda.alignment = centro
        celda.border = borde

    for fila_idx, (_, fila) in enumerate(df.iterrows(), start=2):
        relleno = fila_clara if fila_idx % 2 == 0 else fila_blanca
        for col_idx, valor in enumerate(fila, start=1):
            celda = ws.cell(row=fila_idx, column=col_idx, value=valor)
            celda.font = fuente_normal
            celda.fill = relleno
            celda.alignment = izq
            celda.border = borde

    for col_idx, nombre_col in enumerate(columnas, start=1):
        max_largo = len(str(nombre_col))
        for valor in df.iloc[:, col_idx - 1]:
            max_largo = max(max_largo, len(str(valor)))
        ws.column_dimensions[get_column_letter(col_idx)].width = max_largo + 4

    ws.freeze_panes = "A2"

    # Guardar con el nombre del original + "_ordenado"
    nombre_base = os.path.splitext(os.path.basename(ruta_entrada))[0]
    ruta_salida = os.path.join(carpeta_salida, nombre_base + "_ordenado.xlsx")
    wb.save(ruta_salida)

    return f"Listo, te ordené el archivo. Lo guardé como {nombre_base}_ordenado.xlsx en la carpeta salida."



# Saludo genérico al arrancar la charla ("hola", "buenas", "qué tal"): si es
# así Y es el primer mensaje, Edith no le tira la charla a Gemini, le
# muestra un menú de opciones en el frontend en vez de obligarlo a adivinar
# qué puede pedirle. Si el primer mensaje ya trae algo puntual ("cómo sale
# la selección"), esto no matchea y sigue el camino normal (CHARLA).
PATRON_SALUDO_GENERICO = re.compile(
    r"^(hola+|ey+|hey+|holis+|buenas|buen[oa]s?\s*(d[ií]as?|tardes|noches)?|"
    r"qu[eé]\s*tal|qu[eé]\s*onda|c[oó]mo\s*(est[aá]s?|and[aá]s?|va|vas?))"
    r"(\s*edith)?$",
    re.IGNORECASE,
)


def es_saludo_generico(texto):
    t = (texto or "").strip().lower()
    if not t:
        return False
    # saca signos de puntuación (pueden ir en cualquier lado: "hola,
    # edith!") para no depender de dónde los ponga
    limpio = re.sub(r"[¡¿!?.,]", "", t)
    limpio = re.sub(r"\s+", " ", limpio).strip()
    if not limpio or len(limpio.split()) > 5:
        return False
    return bool(PATRON_SALUDO_GENERICO.match(limpio))


RESPUESTA_SALUDO_MENU = (
    "¡Hola, Nico! ¿En qué te ayudo hoy? Elegí alguna de estas opciones o "
    "contame directamente lo que necesitás."
)


def procesar_texto(texto, historial="", es_primera=False):
    print("Procesando:", texto)

    if es_primera and es_saludo_generico(texto):
        print("Saludo genérico al inicio de la charla: muestro el menú.")
        return RESPUESTA_SALUDO_MENU, False, True

    respuesta_clasificacion = detectar_intencion(texto, historial)
    llamada = respuesta_clasificacion.function_calls[0] if respuesta_clasificacion.function_calls else None

    if llamada:
        print("Herramienta elegida:", llamada.name, llamada.args)
        frase, hay_grafico = ejecutar_herramienta(llamada)
        return frase, hay_grafico, False

    intencion = (respuesta_clasificacion.text or "").strip().upper()
    print("Intención detectada:", intencion)

    if DEMO and intencion in ("MUSICA", "NETFLIX", "EXCEL"):
        return "Esa función solo está disponible en mi versión personal, no acá en la demo pública. Pero puedo charlar con vos o agendarte un recordatorio, por ejemplo.", False, False

    if intencion == "MUSICA":
        cancion = extraer_cancion(texto)
        print("Canción pedida:", cancion)
        if cancion == "NADA":
            abrir_spotify()
            frase = "Abriendo Spotify."
        else:
            frase = reproducir_cancion(cancion)

    elif intencion == "EXCEL":
        frase = convertir_a_excel()

    elif intencion == "NETFLIX":
        termino = extraer_busqueda_netflix(texto)
        print("Buscando en Netflix:", termino)
        frase = buscar_en_netflix(termino)

    else:
        print("Es una charla, le pregunto a Gemini...")
        frase = conversar(texto, historial)

        # Revisamos si hay eventos próximos para recordar
        proximos = eventos_proximos()
        if proximos:
            recordatorios = []
            for id_evento, descripcion, fecha in proximos:
                recordatorios.append(f"{descripcion} (el {fecha})")
                marcar_recordado(id_evento)
            aviso = " Ah, y te recuerdo que tenés: " + ", ".join(recordatorios) + "."
            frase = frase + aviso

        print("Edith responde:", frase)

    return frase, False, False


def armar_respuesta(frase, hay_grafico=False, mostrar_menu=False):
    generar_voz(frase)
    respuesta = send_file("respuesta.mp3", mimetype="audio/mpeg")
    respuesta.headers["X-Respuesta-Texto"] = quote(frase)
    headers_expuestos = "X-Respuesta-Texto"
    if hay_grafico:
        respuesta.headers["X-Hay-Grafico"] = "true"
        headers_expuestos += ", X-Hay-Grafico"
    if mostrar_menu:
        respuesta.headers["X-Mostrar-Menu"] = "true"
        headers_expuestos += ", X-Mostrar-Menu"
    respuesta.headers["Access-Control-Expose-Headers"] = headers_expuestos
    return respuesta

@app.route("/salud", methods=["GET"])
def salud():
    # Chequeo liviano: no toca la base de datos ni Gemini, solo confirma
    # que el proceso de Flask está levantado. Lo usa el frontend para saber
    # cuándo dejar de mostrar "despertando a Edith" y habilitar el chat, y
    # también para saber si mostrar Spotify/Netflix en la presentación
    # inicial (DEMO ya está cargado en memoria, no hace falta tocar nada).
    return {"estado": "ok", "demo": DEMO}, 200


@app.route("/escuchar", methods=["POST"])
def escuchar():
    if DEMO or modelo is None:
        return armar_respuesta("La entrada por voz no está disponible en esta demo pública, pero escribime en el cuadro de texto y te leo la respuesta igual.")

    audio = request.files["audio"]
    audio.save("entrada.wav")
    resultado = modelo.transcribe("entrada.wav", language="Spanish")
    texto = resultado["text"].strip()
    print("Whisper entendió:", texto)

    frase, hay_grafico, mostrar_menu = procesar_texto(texto)
    return armar_respuesta(frase, hay_grafico, mostrar_menu)


@app.route("/texto", methods=["POST"])
def texto():
    mensaje = request.form["texto"]
    historial = request.form.get("historial", "")
    es_primera = request.form.get("primera", "") == "true"
    print("Texto recibido:", mensaje)

    frase, hay_grafico, mostrar_menu = procesar_texto(mensaje, historial, es_primera)
    return armar_respuesta(frase, hay_grafico, mostrar_menu)


@app.route("/grafico", methods=["GET"])
def grafico():
    if not os.path.exists("grafico_gastos.png"):
        return {"error": "Todavía no se generó ningún gráfico."}, 404
    return send_file("grafico_gastos.png", mimetype="image/png")

@app.route("/guardar_conversacion", methods=["POST"])
def endpoint_guardar_conversacion():
    titulo = request.form["titulo"]
    mensajes = request.form["mensajes"]
    print("Guardando conversación:", titulo)
    guardar_conversacion(titulo, mensajes)
    return {"estado": "ok"}

@app.route("/conversaciones", methods=["GET"])
def endpoint_conversaciones():
    filas = listar_conversaciones()
    lista = []
    for id_conv, titulo, fecha in filas:
        lista.append({"id": id_conv, "titulo": titulo, "fecha": str(fecha)})
    return {"conversaciones": lista}

@app.route("/conversacion/<int:id_conv>", methods=["GET"])
def endpoint_obtener_conversacion(id_conv):
    mensajes = obtener_conversacion(id_conv)
    if mensajes:
        return {"mensajes": mensajes}
    return {"mensajes": None}

@app.route("/")
def servir_index():
    return app.send_static_file("index.html")

if __name__ == "__main__":
    puerto = int(os.getenv("PORT", 5001))
    app.run(host="0.0.0.0", port=puerto)
